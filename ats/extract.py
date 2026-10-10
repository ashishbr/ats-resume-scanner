"""Resume text extraction with several independent engines.

Running more than one parser on the same file is how we estimate extraction
accuracy without a ground truth: if independent parsers agree on the words,
an ATS will almost certainly read the same words. Where they disagree, the
layout (columns, text boxes, icon fonts, scanned pages) is the risk.
"""
from __future__ import annotations

import io
import re
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}

# Preference order when engines score equally.
ENGINE_PRIORITY = ["pymupdf", "pdfplumber", "pypdf", "python-docx", "docx-xml", "plain-text"]


@dataclass
class EngineOutput:
    name: str
    text: str = ""
    ok: bool = True
    error: str = ""


@dataclass
class ExtractionResult:
    filename: str
    file_type: str
    engines: list[EngineOutput]
    pages: int = 0
    layout_notes: list[str] = field(default_factory=list)
    links: list[str] = field(default_factory=list)
    header_footer_text: str = ""
    primary_engine: str = ""
    flags: set[str] = field(default_factory=set)  # layout risks: scanned, multi_column, tables, textbox, header_footer

    @property
    def usable_engines(self) -> list[EngineOutput]:
        return [e for e in self.engines if e.ok and e.text.strip()]

    def engine(self, name: str) -> EngineOutput | None:
        return next((e for e in self.engines if e.name == name), None)

    @property
    def primary_text(self) -> str:
        e = self.engine(self.primary_engine)
        if e and e.ok:
            return e.text
        usable = self.usable_engines
        return usable[0].text if usable else ""


def _clean_page_join(pages: list[str]) -> str:
    return "\n\n".join(p.strip("\n") for p in pages).strip() + "\n"


# --------------------------------------------------------------------------- PDF

def _import_pymupdf():
    try:
        import pymupdf  # type: ignore
        return pymupdf
    except ImportError:
        import fitz  # type: ignore  # older PyMuPDF releases
        return fitz


def _pdf_pymupdf(data: bytes, result: ExtractionResult) -> EngineOutput:
    try:
        pymupdf = _import_pymupdf()
    except ImportError:
        return EngineOutput("pymupdf", ok=False, error="PyMuPDF not installed")
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
        pages: list[str] = []
        multi_col_pages, image_pages = [], []
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text")
            pages.append(text)
            # Hyperlinks: an ATS reads visible text only, so a URL hidden behind
            # an icon or the word "LinkedIn" is invisible to it.
            for link in page.get_links():
                uri = link.get("uri")
                if uri and uri not in result.links:
                    result.links.append(uri)
            if len(text.strip()) < 30 and page.get_images():
                image_pages.append(i)
            if _looks_multi_column(page):
                multi_col_pages.append(i)
        result.pages = len(doc)
        if image_pages:
            result.flags.add("scanned")
            result.layout_notes.append(
                f"Page(s) {', '.join(map(str, image_pages))} contain images but almost no text. "
                "Likely scanned; an ATS sees these pages as blank unless it runs OCR."
            )
        if multi_col_pages:
            result.flags.add("multi_column")
            result.layout_notes.append(
                f"Multi-column layout detected on page(s) {', '.join(map(str, multi_col_pages))}. "
                "Some ATS read straight across columns and scramble sections."
            )
        return EngineOutput("pymupdf", _clean_page_join(pages))
    except Exception as exc:  # noqa: BLE001
        return EngineOutput("pymupdf", ok=False, error=f"{type(exc).__name__}: {exc}")


def _looks_multi_column(page) -> bool:
    """Side-by-side columns: several real text lines starting mid-page at the same
    height as lines starting at the left margin. Right-aligned dates (short, and
    starting far right) are ignored, so normal single-column resumes pass."""
    width = page.rect.width
    lines = []
    for block in page.get_text("dict").get("blocks", []):
        for line in block.get("lines", []):
            text = "".join(span.get("text", "") for span in line.get("spans", [])).strip()
            if text:
                lines.append((line["bbox"], text))
    if len(lines) < 8:
        return False
    left_x = min(bb[0] for bb, _ in lines)
    left = [bb for bb, _ in lines if bb[0] < left_x + width * 0.08]
    mid = [bb for bb, t in lines if width * 0.25 < bb[0] < width * 0.6 and len(t) >= 25]
    pairs = sum(1 for mb in mid if any(lb[1] < mb[3] and mb[1] < lb[3] and lb[2] < mb[0] for lb in left))
    return pairs >= 4


def _pdf_pdfplumber(data: bytes, result: ExtractionResult) -> EngineOutput:
    try:
        import pdfplumber  # type: ignore
    except ImportError:
        return EngineOutput("pdfplumber", ok=False, error="pdfplumber not installed")
    try:
        pages, table_pages = [], []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for i, page in enumerate(pdf.pages, start=1):
                pages.append(page.extract_text() or "")
                try:
                    tables = [t for t in page.find_tables() if len(t.rows) > 1 and len(t.rows[0].cells) > 1]
                    if tables:
                        table_pages.append(i)
                except Exception:  # noqa: BLE001
                    pass
            if not result.pages:
                result.pages = len(pdf.pages)
        if table_pages:
            result.flags.add("tables")
            result.layout_notes.append(
                f"Table layout detected on page(s) {', '.join(map(str, table_pages))}. "
                "Older ATS parsers often flatten or skip table cells."
            )
        return EngineOutput("pdfplumber", _clean_page_join(pages))
    except Exception as exc:  # noqa: BLE001
        return EngineOutput("pdfplumber", ok=False, error=f"{type(exc).__name__}: {exc}")


def _pdf_pypdf(data: bytes, result: ExtractionResult) -> EngineOutput:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError:
        return EngineOutput("pypdf", ok=False, error="pypdf not installed")
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:  # noqa: BLE001
                return EngineOutput("pypdf", ok=False, error="PDF is password protected")
        pages = [(p.extract_text() or "") for p in reader.pages]
        if not result.pages:
            result.pages = len(reader.pages)
        return EngineOutput("pypdf", _clean_page_join(pages))
    except Exception as exc:  # noqa: BLE001
        return EngineOutput("pypdf", ok=False, error=f"{type(exc).__name__}: {exc}")


# -------------------------------------------------------------------------- DOCX

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"


def _docx_python_docx(data: bytes, result: ExtractionResult) -> EngineOutput:
    """What a typical ATS sees: header, body paragraphs and tables, footer. No text boxes."""
    try:
        from docx import Document  # type: ignore
        from docx.table import Table  # type: ignore
        from docx.text.paragraph import Paragraph  # type: ignore
    except ImportError:
        return EngineOutput("python-docx", ok=False, error="python-docx not installed")
    try:
        doc = Document(io.BytesIO(data))

        def table_lines(table) -> list[str]:
            lines = []
            for row in table.rows:
                cells, seen = [], set()
                for cell in row.cells:
                    t = cell.text.strip()
                    if t and id(cell._tc) not in seen:  # merged cells repeat
                        cells.append(t)
                    seen.add(id(cell._tc))
                if cells:
                    lines.append(" | ".join(cells))
            return lines

        body: list[str] = []
        table_count = 0
        for child in doc.element.body.iterchildren():
            if child.tag == W + "p":
                para = Paragraph(child, doc)
                is_list = child.find(f"{W}pPr/{W}numPr") is not None or (
                    para.style is not None and para.style.name.lower().startswith("list"))
                body.append(("\u2022 " + para.text) if is_list and para.text.strip() else para.text)
            elif child.tag == W + "tbl":
                table_count += 1
                body.extend(table_lines(Table(child, doc)))

        header, footer = [], []
        for section in doc.sections:
            for part, bucket in ((section.header, header), (section.footer, footer)):
                try:
                    for p in part.paragraphs:
                        if p.text.strip() and p.text not in bucket:
                            bucket.append(p.text)
                    for t in part.tables:
                        for line in table_lines(t):
                            if line not in bucket:
                                bucket.append(line)
                except Exception:  # noqa: BLE001
                    pass

        result.header_footer_text = "\n".join(header + footer)
        if header or footer:
            result.flags.add("header_footer")
            result.layout_notes.append(
                "Content found in the Word header/footer. Several ATS skip headers, so keep "
                "name and contact details in the main body."
            )
        if table_count:
            result.flags.add("tables")
            result.layout_notes.append(
                f"{table_count} table(s) in the document body. Older ATS parsers often flatten or skip table cells."
            )
        # Hyperlink targets
        for rel in doc.part.rels.values():
            if "hyperlink" in rel.reltype and rel.target_ref not in result.links:
                result.links.append(rel.target_ref)
        result.pages = result.pages or 1
        text = "\n".join(header + [""] + body + [""] + footer) if (header or footer) else "\n".join(body)
        return EngineOutput("python-docx", re.sub(r"\n{3,}", "\n\n", text).strip() + "\n")
    except Exception as exc:  # noqa: BLE001
        return EngineOutput("python-docx", ok=False, error=f"{type(exc).__name__}: {exc}")


def _docx_raw_xml(data: bytes, result: ExtractionResult) -> EngineOutput:
    """Reads every w:t node straight from the XML, including text boxes and shapes."""
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
        names = zf.namelist()
        parts = sorted(n for n in names if re.match(r"word/header\d*\.xml$", n))
        parts += ["word/document.xml"]
        parts += sorted(n for n in names if re.match(r"word/footer\d*\.xml$", n))
        out: list[str] = []
        textbox_chars = 0

        def walk(el, buf: list[str], in_textbox: bool) -> None:
            nonlocal textbox_chars
            if el.tag == MC_FALLBACK:  # duplicate legacy copy of a text box
                return
            tag = el.tag
            if tag == W + "t" and el.text:
                buf.append(el.text)
                if in_textbox:
                    textbox_chars += len(el.text)
            elif tag in (W + "tab",):
                buf.append("\t")
            elif tag in (W + "br", W + "cr"):
                buf.append("\n")
            inside = in_textbox or tag == W + "txbxContent"
            for child in el:
                walk(child, buf, inside)
            if tag == W + "p":
                buf.append("\n")
            elif tag == W + "tc":
                buf.append(" | ")

        for part in parts:
            if part not in names:
                continue
            root = ET.fromstring(zf.read(part))
            buf: list[str] = []
            walk(root, buf, False)
            out.append("".join(buf))
        if textbox_chars > 20:
            result.flags.add("textbox")
            result.layout_notes.append(
                f"About {textbox_chars} characters sit inside text boxes/shapes. Many ATS ignore "
                "text boxes entirely, so that content may be invisible to them."
            )
        text = "\n".join(out)
        text = re.sub(r"( \| )+\n", "\n", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return EngineOutput("docx-xml", text.strip() + "\n")
    except Exception as exc:  # noqa: BLE001
        return EngineOutput("docx-xml", ok=False, error=f"{type(exc).__name__}: {exc}")


# --------------------------------------------------------------------------- TXT

def _plain_text(data: bytes) -> EngineOutput:
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            if enc == "utf-16" and not data[:2] in (b"\xff\xfe", b"\xfe\xff"):
                continue
            return EngineOutput("plain-text", text)
        except UnicodeDecodeError:
            continue
    return EngineOutput("plain-text", data.decode("utf-8", errors="replace"))


# ------------------------------------------------------------------------ public

def extract(data: bytes, filename: str) -> ExtractionResult:
    ext = Path(filename).suffix.lower()
    if ext == ".doc":
        raise ValueError("Legacy .doc files are not supported. Save it as .docx or PDF and try again.")
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type '{ext}'. Use PDF, DOCX or TXT.")

    if ext == ".pdf":
        result = ExtractionResult(filename, "pdf", [])
        result.engines = [
            _pdf_pymupdf(data, result),
            _pdf_pdfplumber(data, result),
            _pdf_pypdf(data, result),
        ]
    elif ext == ".docx":
        result = ExtractionResult(filename, "docx", [])
        result.engines = [_docx_python_docx(data, result), _docx_raw_xml(data, result)]
    else:
        result = ExtractionResult(filename, "txt", [_plain_text(data)], pages=1)

    if not result.usable_engines:
        errors = "; ".join(f"{e.name}: {e.error or 'no text'}" for e in result.engines)
        result.layout_notes.append(f"No engine could read text from this file ({errors}).")
    return result


def extract_path(path: str | Path) -> ExtractionResult:
    p = Path(path)
    return extract(p.read_bytes(), p.name)
